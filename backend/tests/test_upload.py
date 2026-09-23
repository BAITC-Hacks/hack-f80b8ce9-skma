import io

from docx import Document
from openpyxl import Workbook

from app.services.attachment_service import MAX_FILE_BYTES


def xlsx(rows):
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def docx(tables=(), paragraphs=()):
    document = Document()
    document.add_paragraph("Спецификация к заказу")
    for rows in tables:
        table = document.add_table(rows=0, cols=len(rows[0]))
        for row in rows:
            for cell, value in zip(table.add_row().cells, row, strict=True):
                cell.text = str(value)
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def upload(client, filename, data, session_id="s1", cart_id=None):
    form = {"session_id": session_id}
    if cart_id:
        form["cart_id"] = cart_id
    return client.post("/api/chat/upload", data=form, files={"file": (filename, data)})


CSV = (
    "Артикул;Наименование;Кол-во\n"
    "200300290;;5\n"  # 027230, stock 19
    "200300280;;20\n"  # 027105, stock 13 -> partial
    "200300285;;2\n"  # 027228, stock 0 -> analog
    ";LED лампа Standart 10W E27;10\n"  # by name
    ";Кабель сверхпроводящий XYZ;1\n"  # not in catalog
)


def by_row(body):
    return {line["row"]: line for line in body["spec"]}


def test_csv_statuses(client):
    response = upload(client, "spec.csv", CSV.encode())
    assert response.status_code == 200, response.text
    lines = by_row(response.json())
    assert lines[2]["status"] == "in_stock" and lines[2]["product"]["id"] == 2
    assert lines[3]["status"] == "partial" and lines[3]["requested_qty"] == 20
    assert lines[4]["status"] == "out_of_stock"
    assert lines[4]["analog"]["id"] == 2 and "160А" in lines[4]["analog"]["reason"]
    assert lines[5]["status"] == "in_stock" and lines[5]["product"]["id"] == 5
    assert lines[6]["status"] == "not_found" and lines[6]["product"] is None


def test_summary_counts_and_cart_untouched(client):
    body = upload(client, "spec.csv", CSV.encode()).json()
    assert "5 позиций" in body["reply"]
    assert "не найдено в каталоге: 1" in body["reply"]
    assert client.get(f"/api/cart/{body['cart_id']}").json()["items"] == []


def test_csv_cp1251_and_comma_delimiter(client):
    data = "Артикул,Количество\n200300290,3\n".encode("cp1251")
    line = upload(client, "spec.csv", data).json()["spec"][0]
    assert line["product"]["id"] == 2 and line["requested_qty"] == 3


def test_xlsx_with_title_row_keeps_file_row_numbers(client):
    data = xlsx(
        [
            ["Спецификация к заказу"],
            [],
            ["№", "Код товара", "Наименование", "Количество, шт"],
            [1, "200300290", "Автомат", 4],
            [2, "027105", None, 1],
        ]
    )
    lines = by_row(upload(client, "spec.xlsx", data).json())
    assert lines[4]["product"]["id"] == 2 and lines[4]["requested_qty"] == 4
    assert lines[5]["product"]["id"] == 3


def test_file_without_header(client):
    data = b"200300290;7\n200300280;1\n"
    lines = upload(client, "spec.csv", data).json()["spec"]
    assert [x["product"]["id"] for x in lines] == [2, 3]
    assert lines[0]["requested_qty"] == 7


def test_docx_takes_spec_table_not_details_table(client):
    data = docx(
        tables=[
            [["Покупатель", "ТОО «Пример»"], ["Контакт", "Отдел закупа"]],
            [["№", "Артикул", "Наименование", "Кол-во"], [1, "200300290", "Автомат", 6],
             [2, "", "LED лампа Standart 10W E27", 3]],
        ]
    )  # fmt: skip
    lines = upload(client, "spec.docx", data).json()["spec"]
    assert [(x["product"]["id"], x["requested_qty"]) for x in lines] == [(2, 6), (5, 3)]


def test_docx_without_tables_reads_lines(client):
    data = docx(paragraphs=["Артикул;Количество", "200300290;2", "200300280;1"])
    lines = upload(client, "spec.docx", data).json()["spec"]
    assert [x["product"]["id"] for x in lines] == [2, 3]


def test_old_doc_format_415(client):
    response = upload(client, "spec.doc", b"\xd0\xcf\x11\xe0")
    assert response.status_code == 415
    assert ".docx" in response.json()["detail"]


def test_broken_docx_422(client):
    assert upload(client, "spec.docx", b"not a zip").status_code == 422


def test_unsupported_type_415(client):
    assert upload(client, "spec.pdf", b"%PDF-1.4").status_code == 415


def test_empty_file_422(client):
    assert upload(client, "spec.csv", b"").status_code == 422


def test_broken_xlsx_422(client):
    assert upload(client, "spec.xlsx", b"not a zip").status_code == 422


def test_too_large_413(client):
    assert upload(client, "spec.csv", b"a" * (MAX_FILE_BYTES + 1)).status_code == 413


def test_llm_sees_uploaded_spec(client):
    from tests.test_chat import FakeLLM, llm_message, use_llm

    fake = FakeLLM(llm_message("Добавить позицию из файла?"))
    use_llm(fake)
    body = upload(client, "spec.csv", CSV.encode())
    client.post(
        "/api/chat",
        json={
            "session_id": "s1",
            "message": "что из файла есть?",
            "cart_id": body.json()["cart_id"],
        },
    )
    system_prompt = fake.requests[0][0].content
    assert "spec.csv" in system_prompt and "product_id=2" in system_prompt
